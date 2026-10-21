"""Admin-initiated external provider import; deduplication by provider ID, never by URL."""
import csv
import hashlib
import json
from pathlib import Path
from datetime import datetime
from sqlalchemy import select
from .db import SessionLocal
from .models import SocialListeningRun, SocialListeningResult, CampaignEvidence

def iter_rows(path):
    path=Path(path)
    if path.suffix.lower()=='.xlsx':
        from openpyxl import load_workbook
        workbook=load_workbook(path,read_only=True,data_only=True)
        try:
            sheet=workbook['all'] if 'all' in workbook.sheetnames else workbook.active
            values=sheet.iter_rows(values_only=True)
            headers=[str(c or '').strip() for c in next(values)]
            for row in values:
                if any(v is not None and str(v).strip() for v in row):
                    yield {key:row[i] if i<len(row) else None for i,key in enumerate(headers)}
        finally:workbook.close()
    else:
        with path.open('r',encoding='utf-8-sig',newline='') as stream:
            yield from csv.DictReader(stream)

def clean(value):
    return str(value).strip() if value is not None else ''

def source_url(row):
    return clean(row.get('Direct URL') or row.get('Reply comment URL') or row.get('Comment URL') or row.get('Post URL') or row.get('Video URL') or row.get('Topic URL'))

def detect_platform(row):
    from urllib.parse import urlparse
    declared=clean(row.get('Source') or row.get('Platform')).lower()
    aliases={'twitter':'x','twitter.com':'x','x.com':'x','fb':'facebook','yt':'youtube'}
    if declared in ('facebook','instagram','youtube','tiktok','threads','x'):
        return declared
    if declared in aliases:return aliases[declared]
    host=urlparse(source_url(row)).hostname or ''
    host=host.lower().removeprefix('www.').removeprefix('m.')
    sites={'facebook.com':'facebook','fb.watch':'facebook','instagram.com':'instagram',
           'x.com':'x','twitter.com':'x','tiktok.com':'tiktok',
           'youtube.com':'youtube','youtu.be':'youtube','threads.net':'threads'}
    return next((name for domain,name in sites.items() if host==domain or host.endswith('.'+domain)),'')

def parse_metric(value):
    import re
    raw=clean(value).replace(',','').replace(' ','').lower()
    if not raw or raw in ('-','n/a','none','nan'):return None
    match=re.fullmatch(r'(\\d+(?:\\.\\d+)?)([kmb]?)',raw)
    if not match:return None
    factor={'':1,'k':1000,'m':1000000,'b':1000000000}[match.group(2)]
    value=int(float(match.group(1))*factor)
    return value if value<10**15 else None

def provider_metrics(row):
    result={}
    for output,keys in {'likes':['Like','Likes'],'comments':['Comment','Comments'],
                        'shares':['Share','Shares'],'views':['View','Views','Play','Plays']}.items():
        for key in keys:
            value=parse_metric(row.get(key))
            if value is not None:
                result[output]=value
                break
    total=parse_metric(row.get('Engagement'))
    if total is not None:result['engagement_total']=total
    elif any(k in result for k in ('likes','comments','shares')):
        result['engagement_total']=sum(result.get(k,0) for k in ('likes','comments','shares'))
    return result

def parse_row(row):
    platform=detect_platform(row)
    ident=clean(row.get('_id'))
    text=clean(row.get('Message'))
    url=source_url(row)
    parent=clean(row.get('Post URL'))
    comment=bool(clean(row.get('Comment URL')) or clean(row.get('Reply comment URL')))
    if not ident or not platform or not url or not text:return None
    return {'platform':platform,'provider_result_id':ident,'text_content':text,'canonical_url':url,
        'content_type':'COMMENT' if comment else 'POST','parent_url':parent if comment else '',
        'published_at':clean(row.get('Post time'))[:64],
        'author_name':clean(row.get('Account'))[:255],
        'engagement':provider_metrics(row),
        'sentiment':clean(row.get('Sentiment')),
        'raw':{str(k):clean(v)[:20000] for k,v in row.items() if k is not None}}

def run_import(run_id,campaign_id,source_path):
    """Background worker. Commits bounded batches and records progress on the run."""
    path=Path(source_path)
    db=SessionLocal()
    run=db.get(SocialListeningRun,run_id)
    if not run:
        db.close();path.unlink(missing_ok=True);return
    count=0;skipped=0;linked=0;processed=0;existing_count=0
    try:
        prior=db.execute(select(SocialListeningResult.platform,SocialListeningResult.provider_result_id,SocialListeningResult.id).where(SocialListeningResult.provider=='ZOCIALEYE')).all()
        known={(platform,pid):rid for platform,pid,rid in prior if pid}
        already=set(db.execute(select(CampaignEvidence.result_id).where(CampaignEvidence.campaign_id==campaign_id)).scalars())
        unique=set()
        run.status='RUNNING';db.commit()
        for number,row in enumerate(iter_rows(path),1):
            processed=number
            item=parse_row(row)
            if not item:
                skipped+=1;continue
            key=(item['platform'],item['provider_result_id'])
            if key in unique:
                skipped+=1;continue
            unique.add(key)
            rid=known.get(key)
            if rid:existing_count+=1
            if not rid:
                raw=item['raw'].copy()
                raw['_enzee_external_import']={'provider':'ZOCIALEYE','parent_url':item['parent_url'],'import_run_id':run_id}
                result=SocialListeningResult(run_id=run_id,provider='ZOCIALEYE',
                    provider_result_id=item['provider_result_id'],platform=item['platform'],
                    content_type=item['content_type'],canonical_url=item['canonical_url'],
                    author_name=item['author_name'],published_at=item['published_at'],
                    text_content=item['text_content'],engagement_json=json.dumps(item['engagement']),
                    source_identity_json=json.dumps({'external_provider':'ZOCIALEYE','parent_url':item['parent_url']}),
                    raw_payload_json=json.dumps(raw,ensure_ascii=False),
                    content_fingerprint=hashlib.sha256((item['platform']+'|'+item['provider_result_id']).encode()).hexdigest())
                db.add(result);db.flush();rid=result.id
                known[key]=rid;count+=1
            if rid not in already:
                db.add(CampaignEvidence(campaign_id=campaign_id,result_id=rid,
                    match_method='EXTERNAL_PROVIDER_IMPORT',matched_fields_json='["admin_requested"]'))
                already.add(rid);linked+=1
            if number%250==0:
                run.result_count=count
                run.query_json=json.dumps({'campaign_id':campaign_id,'processed':number,'inserted':count,'linked':linked,'skipped':skipped,'already_present':existing_count,'provider':'ZOCIALEYE'})
                db.commit()
        run.result_count=count;run.status='COMPLETED'
        run.query_json=json.dumps({'campaign_id':campaign_id,'processed':processed,'inserted':count,'linked':linked,'skipped':skipped,'already_present':existing_count,'provider':'ZOCIALEYE'})
        db.commit()
    except Exception as exc:
        db.rollback()
        run=db.get(SocialListeningRun,run_id)
        if run:
            run.status='FAILED';run.error_text=str(exc)[:800];run.query_json=json.dumps({'campaign_id':campaign_id,'processed':processed,'inserted':count,'linked':linked,'skipped':skipped,'already_present':existing_count,'provider':'ZOCIALEYE','partial_commit':True});db.commit()
    finally:
        db.close();path.unlink(missing_ok=True)
