"""Preview external ZocialEye exports without importing any records."""
import csv
from pathlib import Path

MAPPING={'text_content':('Message',),'canonical_url':('Direct URL','Post URL','Video URL','Topic URL'),
         'author_name':('Account',),'platform':('Source',),'published_at':('Post time',),
         'engagement':('Engagement',),'sentiment':('Sentiment',),'provider_result_id':('_id',),
         'content_type':('Type',),'main_keyword':('Main keyword',),'secondary_keyword':('Sub keyword',)}
def map_fields(headers):
    existing={str(h or '').strip().casefold():str(h).strip() for h in headers if h}
    return {key:next((existing[alias.casefold()] for alias in aliases if alias.casefold() in existing),None) for key,aliases in MAPPING.items()}
def preview_workbook(path):
    path=Path(path)
    if path.suffix.lower()=='.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(path,read_only=True,data_only=True)
        sheets=[]
        for sheet in book:
            headers=[str(v or '').strip() for v in next(sheet.iter_rows(min_row=1,max_row=1,values_only=True),())]
            sheets.append({'name':sheet.title,'data_rows':max(0,(sheet.max_row or 1)-1),'columns':headers})
        names={s['name'] for s in sheets}
        recommended='all' if 'all' in names else next((s['name'] for s in sheets if s['data_rows'] and 'Message' in s['columns']),None)
        if not recommended:raise ValueError('No usable message sheet found')
        ws=book[recommended]
        rows=ws.iter_rows(min_row=2,max_row=min(ws.max_row or 1,26),values_only=True)
        mapping=map_fields(next(s['columns'] for s in sheets if s['name']==recommended))
        keys=next(s['columns'] for s in sheets if s['name']==recommended)
        sampled=list(rows)
        book.close()
    elif path.suffix.lower()=='.csv':
        with path.open('r',encoding='utf-8-sig',newline='') as handle:
            reader=csv.DictReader(handle)
            keys=reader.fieldnames or []
            sampled=[tuple(row.get(key) for key in keys) for _,row in zip(range(25),reader)]
        sheets=[{'name':'CSV','data_rows':sum(1 for _ in path.open('r',encoding='utf-8-sig'))-1,'columns':keys}]
        recommended='CSV';mapping=map_fields(keys)
    else:raise ValueError('Unsupported file type')
    def empty(idx,value):
        return not str(value[idx] if idx<len(value) else '').strip() if idx>=0 else True
    url_key=mapping.get('canonical_url');text_key=mapping.get('text_content')
    ui=keys.index(url_key) if url_key in keys else -1
    ti=keys.index(text_key) if text_key in keys else -1
    return {'sheets':sheets,'recommended_sheet':recommended,
        'estimated_rows_in_recommended_sheet':next(s['data_rows'] for s in sheets if s['name']==recommended),
        'mapped_fields':{k:v for k,v in mapping.items() if v},
        'sampled_rows':len(sampled),'sample_missing_url':sum(empty(ui,row) for row in sampled),
        'sample_missing_text':sum(empty(ti,row) for row in sampled),
        'warning':('Combined all sheet overlaps with platform-specific sheets: importing both would duplicate records. '
                   'Post and comment rows must be preserved separately using source IDs and parent links.') if recommended=='all'
                  else 'Check row types, source URLs, and IDs before enabling import.'}
