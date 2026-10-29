"""Idempotent import of an officer-labelled BAWASLU source workbook.

All rows remain captured campaign evidence; ONLY relevant rows appear in the
existing Intelligence Inbox. Never infer a potential legal violation from a
topic/category label, and never bill or call AI during import.
"""
import hashlib
import io
import json
from collections import Counter
from datetime import datetime
from urllib.parse import urlparse
from openpyxl import load_workbook
from .models import SocialListeningRun, SocialListeningResult, CampaignEvidence, EvidenceAssessment

PROVIDER='HUMAN_VERIFIED_WORKBOOK'
MODEL='HUMAN_VERIFIED_SPREADSHEET'
REQUIRED=('Sample ID','Platform','Account','Message','Direct URL','Relevancy')
MAX_BYTES=12*1024*1024
MAX_ROWS=10000

def read_rows(content):
    if not content or len(content)>MAX_BYTES:
        raise ValueError('Workbook must be nonempty and at most 12 MB')
    try:
        workbook=load_workbook(io.BytesIO(content),read_only=True,data_only=True)
    except Exception as exc:
        raise ValueError('Invalid Excel workbook') from exc
    try:
        sheet=None
        for candidate in workbook.worksheets:
            candidate.reset_dimensions()
            headers=next(candidate.iter_rows(values_only=True),None)
            if headers and all(name in headers for name in REQUIRED):
                sheet=candidate
                break
        if sheet is None:
            raise ValueError('Workbook needs Sample ID, Platform, Account, Message, Direct URL and Relevancy columns')
        iterator=sheet.iter_rows(values_only=True)
        headers=next(iterator)
        columns={str(name).strip():index for index,name in enumerate(headers) if name is not None}
        if not all(k in columns for k in REQUIRED):
            raise ValueError('Incomplete human review column mapping')
        rows=[];seen=set()
        for lineno,values in enumerate(iterator,start=2):
            if all(value is None for value in values):continue
            def v(key):
                idx=columns.get(key)
                val=values[idx] if idx is not None and idx<len(values) else None
                return str(val).strip() if val is not None else ''
            sample=v('Sample ID')
            label=v('Relevancy').upper()
            if not sample or len(sample)>180 or sample in seen:
                raise ValueError('Invalid or repeated Sample ID at row '+str(lineno))
            if label not in ('RELEVANT','IRRELEVANT'):
                raise ValueError('Expected Relevant or Irrelevant at row '+str(lineno))
            url=v('Direct URL') or v('Post URL')
            parsed=urlparse(url)
            if parsed.scheme not in ('http','https') or not parsed.netloc:
                raise ValueError('Missing or invalid source URL at row '+str(lineno))
            message=v('Message')
            if not message:raise ValueError('Missing source message at row '+str(lineno))
            platform=v('Platform').lower()
            if not platform or len(platform)>95:raise ValueError('Invalid platform at row '+str(lineno))
            sentiment=v('Sentiment').upper()
            if sentiment not in ('POSITIVE','NEUTRAL','NEGATIVE'):
                sentiment='UNCLASSIFIED'
            engagement=v('Engagement')
            try:engagement_total=max(0,int(float(engagement))) if engagement else 0
            except (ValueError,TypeError):engagement_total=0
            rows.append({'sample_id':sample,'platform':platform,'account':v('Account')[:255],
                'message':message,'url':url,'published_at':v('Post time')[:64],
                'label':label,'sentiment':sentiment,'category':v('Category'),
                'engagement':engagement_total,'original_row':v('Original Excel Row'),
                'source_id':v('_id'),'main_keyword':v('Main keyword'),
                'sub_keyword':v('Sub keyword'),'image_url':v('Image URL')})
            seen.add(sample)
            if len(rows)>MAX_ROWS:raise ValueError('More than 10,000 records; split workbook')
        if not rows:raise ValueError('No reviewed records found')
        return sheet.title,rows
    finally:
        workbook.close()

def preview(content):
    sheet,rows=read_rows(content)
    counts=Counter(row['label'] for row in rows)
    return {'sheet':sheet,'total':len(rows),'relevant':counts['RELEVANT'],
        'irrelevant':counts['IRRELEVANT'],'unique_direct_urls':len({r['url'] for r in rows}),
        'duplicates_by_sample_id':0,'destination':'ALL_TO_SOCIAL_LISTENING_RELEVANT_ONLY_TO_INBOX',
        'human_reviewed':True,'ai_requests':0,'socialcrawl_credits':0}

def import_rows(db,campaign,content,actor_id):
    sheet,rows=read_rows(content)
    run=SocialListeningRun(provider=PROVIDER,provider_request_id='human-'+hashlib.sha256(content).hexdigest()[:24],
        query_json=json.dumps({'campaign_id':campaign.id,'sheet':sheet,'method':'HUMAN_REVIEWED_EXCEL'}),
        status='COMPLETED',credits_used=0,credits_remaining=0,
        result_count=len(rows),executed_by=str(actor_id))
    db.add(run)
    db.flush()
    sample_ids=['QA:'+r['sample_id'] for r in rows]
    existing={r.provider_result_id:r for r in db.query(SocialListeningResult).filter(
        SocialListeningResult.provider==PROVIDER,
        SocialListeningResult.provider_result_id.in_(sample_ids)).all()}
    inserted=linked=assessed=0
    for row in rows:
        identity='QA:'+row['sample_id']
        result=existing.get(identity)
        if result is None:
            identity_json={'provider':PROVIDER,'sample_id':row['sample_id'],
                'original_source_id':row['source_id']}
            raw={'provenance':'uploaded_human_reviewed_workbook','sample_id':row['sample_id'],
                'original_excel_row':row['original_row'],'human_relevance':row['label'],
                'source_category':row['category'],'main_keyword':row['main_keyword'],
                'sub_keyword':row['sub_keyword'],'image_url':row['image_url']}
            result=SocialListeningResult(run_id=run.id,provider=PROVIDER,provider_result_id=identity,
                platform=row['platform'],content_type='post',canonical_url=row['url'],
                author_name=row['account'],observed_handle=row['account'],
                published_at=row['published_at'],text_content=row['message'],
                relevance_score='',engagement_json=json.dumps({'engagement_total':row['engagement']}),
                source_identity_json=json.dumps(identity_json),raw_payload_json=json.dumps(raw,ensure_ascii=False),
                content_fingerprint=hashlib.sha256((row['platform']+'|'+row['url']+'|'+row['message']).encode()).hexdigest(),
                review_state='HUMAN_REVIEWED',review_notes='Spreadsheet relevance: '+row['label'])
            db.add(result);db.flush()
            inserted+=1
        link=db.query(CampaignEvidence.id).filter_by(campaign_id=campaign.id,result_id=result.id).first()
        if link is None:
            db.add(CampaignEvidence(campaign_id=campaign.id,result_id=result.id,
                match_method='HUMAN_REVIEWED_IMPORT',
                matched_fields_json=json.dumps(['Sample ID','Relevancy'])))
            linked+=1
        latest=(db.query(EvidenceAssessment).filter_by(campaign_id=campaign.id,result_id=result.id)
                .order_by(EvidenceAssessment.assessment_version.desc(),EvidenceAssessment.id.desc()).first())
        if latest and latest.model_version==MODEL:
            if str(latest.relevance_label).upper()!=row['label']:
                raise ValueError('Existing human label changed for '+row['sample_id']+'; manual reconciliation required')
            continue
        db.add(EvidenceAssessment(campaign_id=campaign.id,result_id=result.id,
            assessment_version=latest.assessment_version+1 if latest else 1,
            model_version=MODEL,relevance_label=row['label'],relevance_confidence='HUMAN_REVIEWED',
            sentiment=row['sentiment'],sentiment_confidence='UPLOADED_SOURCE',
            sentiment_explanation='Source workbook sentiment; not independently AI assessed.',
            sentiment_model='UPLOADED_SOURCE_LABEL',
            potential_issue_signals_json='[]',
            topic_labels_json=json.dumps([row['category']] if row['label']=='RELEVANT' and row['category'] not in ('','-') else []),
            explanation='Human-reviewed spreadsheet relevance; source category is not a validated violation.'))
        assessed+=1
    db.commit()
    return {'total':len(rows),'inserted':inserted,'already_present':len(rows)-inserted,
        'linked_to_campaign':linked,'saved_human_assessments':assessed,
        'relevant':sum(r['label']=='RELEVANT' for r in rows),
        'irrelevant':sum(r['label']=='IRRELEVANT' for r in rows),
        'ai_requests':0,'socialcrawl_credits':0,
        'destination':'ALL_SOCIAL_LISTENING_RELEVANT_ONLY_INBOX'}
